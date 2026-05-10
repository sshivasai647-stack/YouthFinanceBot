# Purpose: Platform admin API — user management, analytics, audit logs, exports (Section 4D).
# Existing module dependencies: openpyxl (Excel), pymongo aggregations.

"""Admin routes (role: admin)."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from bson import ObjectId
from bson.errors import InvalidId
from flask import Blueprint, Response, jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity
from marshmallow import Schema, ValidationError, fields, validate
from openpyxl import Workbook

from backend.extensions import get_db
from backend.middleware.role_required import role_required
from backend.services.notification_service import get_user_email, notify_user
from backend.services.serialize import to_jsonable

admin_bp = Blueprint("admin", __name__, url_prefix="/api/admin")


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _admin_oid() -> ObjectId:
    return ObjectId(get_jwt_identity())


class BlockUserSchema(Schema):
    reason = fields.Str(load_default="")


class AssignCounsellorSchema(Schema):
    counsellor_id = fields.Str(required=True)


def _parse_oid(raw: str) -> ObjectId | None:
    try:
        return ObjectId(raw)
    except InvalidId:
        return None


@admin_bp.get("/stats")
@role_required("admin")
def platform_stats():
    db = get_db()
    now = _utc_now()
    week_ago = now - timedelta(days=7)
    users_total = db.users.count_documents({})
    users_active_session = db.users.count_documents({"last_login": {"$gte": week_ago}})
    crisis_flags = db.mental_health_logs.count_documents(
        {"flagged": True, "created_at": {"$gte": now - timedelta(days=1)}}
    )
    chat_today = db.chat_sessions.count_documents({"updated_at": {"$gte": now.replace(hour=0, minute=0, second=0, microsecond=0)}})
    llm_rows = list(
        db.ai_usage.find({"created_at": {"$gte": now.replace(hour=0, minute=0, second=0, microsecond=0)}})
    )
    llm_calls = len(llm_rows)
    avg_ms = (
        sum(int(x.get("latency_ms", 0)) for x in llm_rows) / len(llm_rows) if llm_rows else 0
    )
    return jsonify(
        {
            "users_total": users_total,
            "weekly_active_users": users_active_session,
            "crisis_alerts_last_24h": crisis_flags,
            "chat_sessions_today": chat_today,
            "llm_calls_today": llm_calls,
            "llm_avg_latency_ms": round(avg_ms, 1),
        }
    ), 200


@admin_bp.get("/users")
@role_required("admin")
def list_users():
    db = get_db()
    page = max(int(request.args.get("page", "1")), 1)
    limit = min(max(int(request.args.get("limit", "20")), 1), 100)
    skip = (page - 1) * limit
    role_f = request.args.get("role")
    active_raw = request.args.get("active")
    q: dict[str, Any] = {}
    if role_f:
        q["role"] = role_f
    if active_raw is not None:
        if active_raw.lower() in ("true", "1"):
            q["is_active"] = True
        elif active_raw.lower() in ("false", "0"):
            q["is_active"] = False
    total = db.users.count_documents(q)
    items = []
    for u in db.users.find(q).sort("created_at", -1).skip(skip).limit(limit):
        items.append(
            {
                "id": str(u["_id"]),
                "name": u.get("name", ""),
                "email": u.get("email", ""),
                "role": u.get("role"),
                "age": u.get("age"),
                "is_active": u.get("is_active", True),
                "is_verified": u.get("is_verified", False),
                "counsellor_id": str(u["counsellor_id"]) if u.get("counsellor_id") else None,
                "counselling_status": u.get("counselling_status"),
                "created_at": u.get("created_at"),
                "last_login": u.get("last_login"),
            }
        )
    return jsonify({"total": total, "page": page, "users": to_jsonable(items)}), 200


@admin_bp.patch("/users/<user_id>/block")
@role_required("admin")
def block_user(user_id: str):
    uid = _parse_oid(user_id)
    if not uid or uid == _admin_oid():
        return jsonify({"error": "Invalid user"}), 400
    try:
        body = BlockUserSchema().load(request.get_json() or {})
    except ValidationError as err:
        return jsonify({"error": "Validation failed", "details": err.messages}), 400
    db = get_db()
    u = db.users.find_one_and_update(
        {"_id": uid},
        {"$set": {"is_active": False}},
    )
    if not u:
        return jsonify({"error": "User not found"}), 404
    email = get_user_email(db.users, uid)
    notify_user(
        db.notifications,
        user_id=uid,
        notif_type="account_suspended",
        message="Your YouthFinanceBot account has been suspended. Contact support if this is a mistake.",
        send_email=True,
        email_subject="YouthFinanceBot — Account suspended",
        recipient_email=email,
    )
    db.audit_logs.insert_one(
        {
            "user_id": str(_admin_oid()),
            "action": "block_user",
            "endpoint": f"/api/admin/users/{user_id}/block",
            "ip": request.remote_addr,
            "timestamp": _utc_now(),
            "target_user": str(uid),
            "reason": body.get("reason", ""),
        }
    )
    return jsonify({"message": "User blocked"}), 200


@admin_bp.patch("/users/<user_id>/unblock")
@role_required("admin")
def unblock_user(user_id: str):
    uid = _parse_oid(user_id)
    if not uid:
        return jsonify({"error": "Invalid user"}), 400
    db = get_db()
    res = db.users.update_one({"_id": uid}, {"$set": {"is_active": True}})
    if res.matched_count == 0:
        return jsonify({"error": "User not found"}), 404
    return jsonify({"message": "User unblocked"}), 200


@admin_bp.patch("/users/<user_id>/assign-counsellor")
@role_required("admin")
def assign_counsellor(user_id: str):
    uid = _parse_oid(user_id)
    if not uid:
        return jsonify({"error": "Invalid user"}), 400
    try:
        body = AssignCounsellorSchema().load(request.get_json() or {})
    except ValidationError as err:
        return jsonify({"error": "Validation failed", "details": err.messages}), 400
    cid = _parse_oid(body["counsellor_id"])
    if not cid:
        return jsonify({"error": "Invalid counsellor id"}), 400
    db = get_db()
    counsellor = db.users.find_one({"_id": cid, "role": "counsellor"})
    if not counsellor:
        return jsonify({"error": "Counsellor not found"}), 404
    db.users.update_one({"_id": uid, "role": "citizen"}, {"$set": {"counsellor_id": cid}})
    return jsonify({"message": "Counsellor assigned"}), 200


@admin_bp.get("/analytics/modules")
@role_required("admin")
def analytics_modules():
    db = get_db()
    since = _utc_now() - timedelta(days=30)
    pipeline = [
        {"$match": {"timestamp": {"$gte": since}, "endpoint": {"$regex": "^/api/"}}},
        {
            "$project": {
                "mod": {
                    "$arrayElemAt": [
                        {"$split": ["$endpoint", "/"]},
                        2,
                    ]
                }
            }
        },
        {"$group": {"_id": "$mod", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
    ]
    try:
        rows = list(db.audit_logs.aggregate(pipeline))
    except Exception:
        rows = []
    usage = {r["_id"]: r["count"] for r in rows if r.get("_id")}
    return jsonify({"since_days": 30, "module_hits": usage}), 200


@admin_bp.get("/analytics/trends")
@role_required("admin")
def analytics_trends():
    db = get_db()
    since = _utc_now() - timedelta(days=56)
    weekly: list[dict[str, Any]] = []
    for i in range(8):
        start = since + timedelta(days=i * 7)
        end = start + timedelta(days=7)
        uq = {
            "created_at": {"$gte": start, "$lt": end},
            "last_login": {"$gte": start, "$lt": end},
        }
        # Simplified: count distinct-ish using last_login buckets
        active = db.users.count_documents({"last_login": {"$gte": start, "$lt": end}})
        complaints = db.mental_health_logs.count_documents(
            {"flagged": True, "created_at": {"$gte": start, "$lt": end}}
        )
        weekly.append(
            {
                "week_start": start.date().isoformat(),
                "weekly_active_logins": active,
                "flagged_mental_health_events": complaints,
            }
        )
    return jsonify({"weeks": weekly}), 200


@admin_bp.get("/analytics/ai-usage")
@role_required("admin")
def analytics_ai_usage():
    db = get_db()
    since = _utc_now() - timedelta(days=7)
    rows = list(db.ai_usage.find({"created_at": {"$gte": since}}))
    n = len(rows)
    avg_ms = sum(int(r.get("latency_ms", 0)) for r in rows) / n if n else 0
    err_ct = sum(1 for r in rows if r.get("had_llm_error"))
    return jsonify(
        {
            "calls": n,
            "avg_latency_ms": round(avg_ms, 2),
            "llm_errors": err_ct,
            "since_days": 7,
        }
    ), 200


@admin_bp.get("/logs")
@role_required("admin")
def audit_logs_list():
    db = get_db()
    page = max(int(request.args.get("page", "1")), 1)
    limit = min(max(int(request.args.get("limit", "50")), 1), 200)
    skip = (page - 1) * limit
    q: dict[str, Any] = {}
    ep = request.args.get("endpoint_prefix")
    if ep:
        q["endpoint"] = {"$regex": ep}
    total = db.audit_logs.count_documents(q)
    entries = []
    for doc in db.audit_logs.find(q).sort("timestamp", -1).skip(skip).limit(limit):
        entries.append(
            {
                "id": str(doc["_id"]),
                "user_id": doc.get("user_id"),
                "action": doc.get("action"),
                "endpoint": doc.get("endpoint"),
                "ip": doc.get("ip"),
                "timestamp": doc.get("timestamp"),
            }
        )
    return jsonify({"total": total, "page": page, "logs": to_jsonable(entries)}), 200


@admin_bp.get("/analytics/export")
@role_required("admin")
def analytics_export():
    db = get_db()
    wb = Workbook()
    ws_u = wb.active
    ws_u.title = "Users"
    ws_u.append(["id", "email", "name", "role", "is_active", "created_at", "last_login"])
    for u in db.users.find({}).limit(5000):
        ws_u.append(
            [
                str(u["_id"]),
                u.get("email", ""),
                u.get("name", ""),
                u.get("role", ""),
                u.get("is_active", True),
                str(u.get("created_at", "")),
                str(u.get("last_login", "")),
            ]
        )
    ws_m = wb.create_sheet("AuditSample")
    ws_m.append(["endpoint", "action", "user_id", "timestamp"])
    for a in db.audit_logs.find({}).sort("timestamp", -1).limit(5000):
        ws_m.append(
            [
                a.get("endpoint", ""),
                a.get("action", ""),
                a.get("user_id", ""),
                str(a.get("timestamp", "")),
            ]
        )
    ws_a = wb.create_sheet("AIUsage")
    ws_a.append(["latency_ms", "had_llm_error", "created_at"])
    for r in db.ai_usage.find({}).sort("created_at", -1).limit(5000):
        ws_a.append([r.get("latency_ms"), r.get("had_llm_error"), str(r.get("created_at", ""))])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="youthfinancebot_analytics.xlsx",
    )
